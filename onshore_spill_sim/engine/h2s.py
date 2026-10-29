# -*- coding: utf-8 -*-
"""
Hydrogen sulphide (H2S) leakage at a well - Gaussian puff model.

Source term
  Gas rate G (MMscf/d at 60 F, 14.7 psia) with H2S mole fraction x.
  Molar flow n = G 1e6 0.0283168 / 86400 x P/(R T_std); H2S mass rate q = n x 34.08 g/s.
  The gas leaves the rig floor / wellhead as a vertical jet of diameter d at the
  actual exit velocity v = n R T_gas / (P pi d^2 / 4), capped at 400 m/s (choked flow).

Plume rise (Briggs, 1975, 1984)
  buoyancy flux Fb = g v r^2 (rho_air - rho_gas) / rho_air
  momentum flux Fm = v^2 r^2 rho_gas / rho_air
  neutral / unstable: dh(x) = 1.6 Fb^1/3 x^2/3 / u up to the final rise
                      (21.425 Fb^3/4 / u for Fb < 55, else 38.71 Fb^3/5 / u),
                      momentum rise 3 d v / u, the larger is used
  stable (E, F):      dh = 2.6 (Fb / (u s))^1/3, momentum 1.5 (Fm / (u s^0.5))^1/3,
                      s = (g / T) dtheta/dz
  Gas lighter than air (natural gas with a few per cent H2S) rises; a mixture heavier
  than air is flagged, because a passive Gaussian model is then not conservative
  close to the source.

Dispersion
  The release is split into puffs (one every puff_s seconds).  Each puff is carried by
  the hourly wind at plume height (power-law profile, exponent by stability class),
  turning smoothly between hours, so the plume bends as the wind veers.
  Puff spread follows the Briggs (1973) open-country curves for the
  Pasquill-Gifford class of the moment, never shrinking.  Ground-level concentration:
      C = 2 m / ((2 pi)^1.5 sy^2 sz) exp(-r^2 / 2 sy^2) [exp(-He^2/2sz^2) + lid images]
  with reflection at the ground and at the top of the mixed layer, uniform mixing once
  sz > 1.6 L, and the half-height terrain correction (plume height reduced by half the
  rise of the ground above the wellhead).

Stability class: Turner (1964) scheme from wind speed, solar elevation and cloud.

Concentrations: ppm = mg/m3 x 24.45 / 34.08 (25 C, 1 atm).  P-G spreads represent
about 10-minute averages; 60-minute averages are the running mean of the last hour.

Thresholds (EPA Acute Exposure Guideline Levels for H2S, 60-minute values; NIOSH IDLH):
  AEGL-1 0.51 ppm (discomfort, odour, reversible), AEGL-2 27 ppm (irreversible or
  escape-impairing effects), AEGL-3 50 ppm (life-threatening), IDLH 100 ppm,
  500 ppm (rapid loss of consciousness).
"""
import math
import os
import json
import csv
import datetime as _dt

import numpy as np
from osgeo import gdal, ogr, osr

from .terrain import read_dem_window
from . import heat as ht
from . import climate, fallout, outputs, html_export
from .receptors import load_receptors

R_GAS = 8.314462
P_ATM = 101325.0
MW_H2S = 34.08
MW_AIR = 28.97
PPM_PER_MGM3 = 24.45 / MW_H2S

THRESHOLDS = [
    # key, label, ppm, basis ('1h' = 60-min mean, 'peak' = 10-min mean), colour
    ("AEGL1", "AEGL-1 (60 min) 0.51 ppm", 0.51, "1h", "#2f9e44"),
    ("AEGL2", "AEGL-2 (60 min) 27 ppm", 27.0, "1h", "#f08c00"),
    ("AEGL3", "AEGL-3 (60 min) 50 ppm", 50.0, "1h", "#e03131"),
    ("IDLH", "IDLH 100 ppm", 100.0, "peak", "#9c36b5"),
    ("P500", "500 ppm - rapid loss of consciousness", 500.0, "peak", "#3b0a45"),
]

H2S_RAMP = [  # (ppm, r, g, b, a)
    (0.1, 150, 220, 235, 70),
    (0.51, 80, 190, 120, 130),
    (5.0, 215, 215, 60, 165),
    (27.0, 255, 146, 0, 200),
    (50.0, 235, 40, 30, 220),
    (100.0, 170, 0, 120, 235),
    (500.0, 70, 0, 70, 250),
]

# Briggs (1973) open-country coefficients, classes A..F
_SY = [(0.22, 1e-4), (0.16, 1e-4), (0.11, 1e-4), (0.08, 1e-4), (0.06, 1e-4), (0.04, 1e-4)]
_P_EXP = [0.07, 0.07, 0.10, 0.15, 0.35, 0.55]
_MIX = [1500.0, 1300.0, 1000.0, 800.0, 300.0, 150.0]
_DTHETA = [0.0, 0.0, 0.0, 0.0, 0.02, 0.035]


def _sz_class(x, k):
    if k == 0:
        return 0.20 * x
    if k == 1:
        return 0.12 * x
    if k == 2:
        return 0.08 * x / math.sqrt(1 + 2e-4 * x)
    if k == 3:
        return 0.06 * x / math.sqrt(1 + 1.5e-3 * x)
    if k == 4:
        return 0.03 * x / (1 + 3e-4 * x)
    return 0.016 * x / (1 + 3e-4 * x)


def sigmas(x, cls):
    """sigma_y, sigma_z (m) at travel distance x (m) for stability index cls (1=A .. 6=F,
    fractional values interpolate geometrically)."""
    x = max(x, 1.0)
    k0 = int(min(max(math.floor(cls - 1), 0), 5))
    k1 = min(k0 + 1, 5)
    f = min(max(cls - 1 - k0, 0.0), 1.0)
    sy = [a * x / math.sqrt(1 + b * x) for a, b in (_SY[k0], _SY[k1])]
    sz = [_sz_class(x, k0), _sz_class(x, k1)]
    gy = sy[0] ** (1 - f) * sy[1] ** f
    gz = sz[0] ** (1 - f) * sz[1] ** f
    return gy, min(gz, 5000.0)


def stability_index(sin_elev, cloud, u10):
    """Pasquill-Gifford class (1=A ... 6=F) by the Turner insolation scheme."""
    if cloud >= 0.9:
        return 4.0
    if sin_elev > 0.26:  # sun higher than 15 deg - daytime
        el = math.degrees(math.asin(min(sin_elev, 1.0)))
        ins = 0 if el > 60 else (1 if el > 35 else 2)       # strong, moderate, slight
        if cloud > 0.5:
            ins = min(ins + 1, 2)
        table = [(2.0, (1.0, 1.5, 2.0)), (3.0, (1.5, 2.0, 3.0)), (5.0, (2.0, 2.5, 3.0)),
                 (6.0, (3.0, 3.5, 4.0)), (99.0, (3.0, 4.0, 4.0))]
        for lim, row in table:
            if u10 < lim:
                return row[ins]
    if sin_elev > 0.0:
        return 4.0  # around sunrise / sunset
    cloudy = cloud >= 0.5
    for lim, row in [(2.0, (6.0, 5.0)), (3.0, (6.0, 5.0)), (5.0, (5.0, 4.0)), (99.0, (4.0, 4.0))]:
        if u10 < lim:
            return row[1] if cloudy else row[0]
    return 4.0


def _cls_param(table, cls):
    k0 = int(min(max(math.floor(cls - 1), 0), 5))
    k1 = min(k0 + 1, 5)
    f = min(max(cls - 1 - k0, 0.0), 1.0)
    return table[k0] * (1 - f) + table[k1] * f


def plume_rise(Fb, Fm, d, v, u, cls, x, Ta_k):
    u = max(u, 1.0)
    if cls >= 4.5:   # stable
        s = 9.81 / Ta_k * _cls_param(_DTHETA, cls)
        s = max(s, 1e-4)
        final_b = 2.6 * (max(Fb, 0.0) / (u * s)) ** (1.0 / 3.0) if Fb > 0 else 0.0
        grad_b = 1.6 * max(Fb, 0.0) ** (1.0 / 3.0) * x ** (2.0 / 3.0) / u
        dh_b = min(grad_b, final_b)
        dh_m = min(1.5 * (Fm / (u * math.sqrt(s))) ** (1.0 / 3.0), 3.0 * d * v / u)
    else:
        if Fb > 0:
            if Fb < 55:
                final_b = 21.425 * Fb ** 0.75 / u
            else:
                final_b = 38.71 * Fb ** 0.6 / u
            grad_b = 1.6 * Fb ** (1.0 / 3.0) * x ** (2.0 / 3.0) / u
            dh_b = min(grad_b, final_b)
        else:
            dh_b = 0.0
        dh_m = 3.0 * d * v / u
    return max(dh_b, dh_m)


class H2SSource(object):
    def __init__(self, gas_mmscfd=10.0, h2s_ppm=10000.0, exit_d_m=0.20, release_h_m=10.0,
                 gas_temp_c=40.0, gas_mw=19.0, duration_h=24.0, jet=True, buildup_h=0.0):
        self.buildup_h = float(buildup_h or 0.0)
        self.gas_mmscfd = gas_mmscfd
        self.h2s_ppm = h2s_ppm
        self.exit_d_m = exit_d_m
        self.release_h_m = release_h_m
        self.gas_temp_c = gas_temp_c
        self.gas_mw = gas_mw
        self.duration_h = duration_h
        self.jet = jet
        x = h2s_ppm / 1e6
        q_std = gas_mmscfd * 1e6 * 0.0283168 / 86400.0        # m3/s at 60 F
        self.n_mol = q_std * P_ATM / (R_GAS * 288.71)          # mol/s
        self.q_gs = self.n_mol * x * MW_H2S                    # g/s H2S
        Tg = gas_temp_c + 273.15
        self.mw_mix = x * MW_H2S + (1 - x) * gas_mw
        self.rho_gas = P_ATM * self.mw_mix / 1000.0 / (R_GAS * Tg)
        q_act = self.n_mol * R_GAS * Tg / P_ATM
        area = math.pi * (exit_d_m / 2.0) ** 2
        self.v_exit = min(q_act / max(area, 1e-6), 400.0) if jet else 1.0
        self.q_act = q_act

    def rate_factor(self, t_s):
        """Share of the full leak rate at time t (s): a leak that builds up starts at 10 % and reaches the
        full rate after buildup_h hours (a leak at full rate from the start when buildup_h = 0)."""
        if self.buildup_h <= 0:
            return 1.0
        return 0.1 + 0.9 * min(1.0, t_s / (self.buildup_h * 3600.0))

    def fluxes(self, Ta_c):
        Ta = Ta_c + 273.15
        rho_a = P_ATM * MW_AIR / 1000.0 / (R_GAS * Ta)
        r = self.exit_d_m / 2.0
        v = self.v_exit
        Fb = 9.81 * v * r * r * (rho_a - self.rho_gas) / rho_a
        Fm = v * v * r * r * self.rho_gas / rho_a
        return Fb, Fm, rho_a

    def dense(self, Ta_c):
        return self.fluxes(Ta_c)[0] < 0


def _interp_wind(wind, t_h):
    """Wind (to_deg, speed) interpolated between hourly values."""
    i = int(math.floor(t_h))
    i0 = min(max(i, 0), len(wind) - 1)
    i1 = min(i0 + 1, len(wind) - 1)
    f = t_h - i
    a0, a1 = wind[i0][0], wind[i1][0]
    d = ((a1 - a0 + 180.0) % 360.0) - 180.0
    return (a0 + f * d) % 360.0, wind[i0][1] * (1 - f) + wind[i1][1] * f


def simulate(dem, source_xy, src, wind, month, start_hour, doy, sim_hours, frame_minutes=10.0,
             half_width_m=10000.0, cell_m=50.0, puff_s=90.0, lat=9.3, lon=-0.5, cloud=0.35,
             t_mean=28.0, t_max=34.0, u_min=1.0, log=print, feedback=None):
    """Run the puff model; returns a dict of arrays and per-frame records."""
    grid = read_dem_window(dem, source_xy[0], source_xy[1], half_width_m, target_cell=cell_m)
    nr, nc = grid.nrows, grid.ncols
    dx = grid.dx
    X = grid.gt[0] + (np.arange(nc) + 0.5) * grid.gt[1]
    Y = grid.gt[3] + (np.arange(nr) + 0.5) * grid.gt[5]
    sr, scc = grid.xy_to_rc(*source_xy)
    zs = float(grid.z[sr, scc])
    zg = grid.z
    dt_eval = frame_minutes * 60.0
    n_eval = int(round(sim_hours * 3600.0 / dt_eval)) + 1
    per_hour = max(int(round(3600.0 / dt_eval)), 1)
    puffs = []           # dict(x, y, dist, sy, sz, m, t0)
    peak = np.zeros((nr, nc), np.float32)
    max1h = np.zeros((nr, nc), np.float32)
    t_aegl2 = np.full((nr, nc), np.nan, np.float32)
    hist = []
    frames = []
    stats = []
    t = 0.0
    next_puff = 0.0
    step = puff_s / math.ceil(puff_s / 60.0)      # sub-step that divides the puff interval
    dense_flag = src.dense(t_mean)
    log("H2S source: %.1f MMscf/d gas, %.0f ppm H2S -> %.2f g/s (%.1f kg/h); exit %.0f m/s; %s." % (
        src.gas_mmscfd, src.h2s_ppm, src.q_gs, src.q_gs * 3.6, src.v_exit,
        "mixture heavier than air - dense-gas behaviour near the source is not represented"
        if dense_flag else "mixture lighter than air (buoyant plume)"))
    for k in range(n_eval):
        t_target = k * dt_eval
        # advance puffs to t_target in sub-steps
        while t < t_target - 1e-6:
            dts = min(step, t_target - t)
            th = t / 3600.0
            hod = (start_hour + th) % 24.0
            day = doy + int((start_hour + th) // 24)
            se = ht.solar_elevation_sin(lat, day, hod, lon)
            to_deg, u10 = _interp_wind(wind, th)
            u10 = max(u10, u_min)
            cls = stability_index(se, cloud, u10)
            Ta = ht.air_temperature(t_mean, t_max, hod)
            # release
            while next_puff <= t + 1e-6 and next_puff < src.duration_h * 3600.0:
                puffs.append(dict(x=source_xy[0], y=source_xy[1], dist=0.0, sy=1.0, sz=1.0,
                                  m=src.q_gs * puff_s * src.rate_factor(next_puff), H=src.release_h_m, age=0.0))
                next_puff += puff_s
            p_exp = _cls_param(_P_EXP, cls)
            Fb, Fm, _ = src.fluxes(Ta)
            ang = math.radians(to_deg)
            for p in puffs:
                hp = max(p["H"], 10.0)
                u = u10 * (min(hp, 300.0) / 10.0) ** p_exp
                d = u * dts
                p["x"] += d * math.sin(ang)
                p["y"] += d * math.cos(ang)
                p["dist"] += d
                sy, sz = sigmas(p["dist"], cls)
                p["sy"] = max(p["sy"], sy)
                p["sz"] = max(p["sz"], sz)
                if src.jet:
                    dh = plume_rise(Fb, Fm, src.exit_d_m, src.v_exit, u, cls, max(p["dist"], 1.0),
                                    Ta + 273.15)
                else:
                    dh = 0.0
                p["H"] = max(p["H"], src.release_h_m + dh)
                p["cls"] = cls
                p["age"] += dts
            # drop puffs that have left the domain
            x0, y0, x1, y1 = grid.extent()
            puffs = [p for p in puffs if (x0 - 3 * p["sy"] < p["x"] < x1 + 3 * p["sy"] and
                                          y0 - 3 * p["sy"] < p["y"] < y1 + 3 * p["sy"])]
            t += dts
        # ground-level concentration at t_target
        th = t_target / 3600.0
        hod = (start_hour + th) % 24.0
        day = doy + int((start_hour + th) // 24)
        se = ht.solar_elevation_sin(lat, day, hod, lon)
        to_deg, u10 = _interp_wind(wind, th)
        u10 = max(u10, u_min)
        cls = stability_index(se, cloud, u10)
        L = _cls_param(_MIX, cls)
        C = np.zeros((nr, nc), np.float64)  # g/m3
        # slug sampling: each puff is spread along the segment to the next older puff
        # so that the plume is continuous even where the puff spacing exceeds sigma_y
        samples = []
        for i, p in enumerate(puffs):
            q = puffs[i - 1] if i > 0 else None
            if q is not None and q["age"] - p["age"] <= 2.5 * puff_s:
                seg = math.hypot(q["x"] - p["x"], q["y"] - p["y"])
                n = int(min(max(math.ceil(seg / max(0.5 * p["sy"], 1.0)), 1), 40))
            else:
                n = 1
            for j in range(n):
                f = (j + 0.5) / n if n > 1 else 0.0
                if n > 1:
                    samples.append((p["x"] + f * (q["x"] - p["x"]), p["y"] + f * (q["y"] - p["y"]),
                                    p["sy"] + f * (q["sy"] - p["sy"]), p["sz"] + f * (q["sz"] - p["sz"]),
                                    p["H"] + f * (q["H"] - p["H"]), p["m"] / n))
                else:
                    samples.append((p["x"], p["y"], p["sy"], p["sz"], p["H"], p["m"]))
        for (px_, py_, sy, sz, Hp, mp) in samples:
            rad = 4.0 * sy
            c0 = int(max(math.floor((px_ - rad - grid.gt[0]) / dx), 0))
            c1 = int(min(math.ceil((px_ + rad - grid.gt[0]) / dx) + 1, nc))
            r0 = int(max(math.floor((grid.gt[3] - (py_ + rad)) / dx), 0))
            r1 = int(min(math.ceil((grid.gt[3] - (py_ - rad)) / dx) + 1, nr))
            if c1 <= c0 or r1 <= r0:
                continue
            ddx = (X[c0:c1] - px_)[None, :]
            ddy = (Y[r0:r1] - py_)[:, None]
            hor = np.exp(-(ddx ** 2 + ddy ** 2) / (2 * sy * sy))
            # half-height terrain correction: ground above the wellhead lowers the plume
            # by half of the rise; lower ground leaves it unchanged
            rise = np.maximum(zg[r0:r1, c0:c1] - zs, 0.0)
            He = np.minimum(np.maximum(Hp - 0.5 * rise, 0.0), L)
            if sz > 1.6 * L:
                C[r0:r1, c0:c1] += mp * hor / (2 * math.pi * sy * sy * L)
            else:
                ver = np.exp(-He ** 2 / (2 * sz * sz))
                for n_ in (1, -1):
                    ver = ver + np.exp(-(2 * n_ * L - He) ** 2 / (2 * sz * sz))
                C[r0:r1, c0:c1] += 2.0 * mp * hor * ver / ((2 * math.pi) ** 1.5 * sy * sy * sz)
        ppm = (C * 1000.0 * PPM_PER_MGM3).astype(np.float32)
        hist.append(ppm)
        if len(hist) > per_hour:
            hist.pop(0)
        m1 = np.mean(hist, axis=0) if len(hist) == per_hour else np.sum(hist, axis=0) / per_hour
        np.maximum(peak, ppm, out=peak)
        np.maximum(max1h, m1, out=max1h)
        new2 = np.isnan(t_aegl2) & (m1 >= 27.0)
        t_aegl2[new2] = th
        sel = ppm >= 0.05
        rr, cc = np.nonzero(sel)
        frames.append(dict(t_h=th, rows=rr.astype(np.int32), cols=cc.astype(np.int32),
                           ppm=ppm[sel].astype(np.float32),
                           wind=list(wind[min(int(th), len(wind) - 1)]),
                           stab="ABCDEF"[int(round(cls)) - 1], rain_mmh=0.0))
        dd = {}
        for key, lab, thr, basis, col in THRESHOLDS:
            arr = m1 if basis == "1h" else ppm
            m = arr >= thr
            if m.any():
                r_, c_ = np.nonzero(m)
                d = np.hypot(X[c_] - source_xy[0], Y[r_] - source_xy[1])
                dd[key] = float(d.max())
                dd[key + "_ha"] = float(m.sum()) * dx * dx / 1e4
            else:
                dd[key] = 0.0
                dd[key + "_ha"] = 0.0
        stats.append(dict(t_h=th, max_ppm=float(ppm.max()), max_1h_ppm=float(m1.max()),
                          stability="ABCDEF"[int(round(cls)) - 1], wind_from=(to_deg + 180) % 360,
                          wind_mps=u10, mixing_m=L, puffs=len(puffs),
                          plume_h_m=float(np.mean([p["H"] for p in puffs])) if puffs else 0.0,
                          releasing=(src.rate_factor(t_target) if t_target < src.duration_h * 3600.0 else 0.0),
                          **dd))
        if feedback is not None:
            feedback.setProgress(100.0 * k / max(n_eval - 1, 1))
            if feedback.isCanceled():
                break
    return dict(grid=grid, peak=peak, max1h=max1h, t_aegl2=t_aegl2, frames=frames, stats=stats,
                X=X, Y=Y, dense=dense_flag)


# ---------------------------------------------------------------------------
def _hazard_class(peak, m1):
    if peak >= 500:
        return "500 ppm+"
    if peak >= 100:
        return "IDLH"
    if m1 >= 50:
        return "AEGL-3"
    if m1 >= 27:
        return "AEGL-2"
    if m1 >= 0.51:
        return "AEGL-1"
    if peak >= 0.05:
        return "Odour only"
    return "Not reached"


def receptor_exposure(res, recs, source_xy, frame_h):
    g = res["grid"]
    out = []
    shape = (g.nrows, g.ncols)
    per_hour = max(int(round(1.0 / frame_h)), 1)
    for rc in recs:
        r, c = g.xy_to_rc(rc["x"], rc["y"])
        item = dict(rc)
        item.update(dist_km=math.hypot(rc["x"] - source_xy[0], rc["y"] - source_xy[1]) / 1000.0,
                    peak_ppm=0.0, max_1h_ppm=0.0, first_aegl1_h=None, first_aegl2_h=None,
                    first_aegl3_h=None, first_idlh_h=None, inside=g.inside(r, c), series=[])
        if g.inside(r, c):
            item["peak_ppm"] = float(res["peak"][r, c])
            item["max_1h_ppm"] = float(res["max1h"][r, c])
            ser = []
            for fr in res["frames"]:
                hit = np.nonzero((fr["rows"] == r) & (fr["cols"] == c))[0]
                ser.append(float(fr["ppm"][hit[0]]) if hit.size else 0.0)
            item["series"] = ser
            hist = []
            for i, v in enumerate(ser):
                hist.append(v)
                hist = hist[-per_hour:]
                m1 = sum(hist) / per_hour
                th = res["frames"][i]["t_h"]
                for key, thr in (("first_aegl1_h", 0.51), ("first_aegl2_h", 27.0), ("first_aegl3_h", 50.0)):
                    if item[key] is None and m1 >= thr:
                        item[key] = th
                if item["first_idlh_h"] is None and v >= 100.0:
                    item["first_idlh_h"] = th
        item["hazard"] = _hazard_class(item["peak_ppm"], item["max_1h_ppm"])
        out.append(item)
    return out


def _ramp_png_frames(res, folder, prefix, log=print):
    g = res["grid"]
    os.makedirs(folder, exist_ok=True)
    drv = gdal.GetDriverByName("MEM")
    png = gdal.GetDriverByName("PNG")
    paths = []
    bounds = None
    for i, fr in enumerate(res["frames"]):
        a = np.zeros((g.nrows, g.ncols), np.float32)
        a[fr["rows"], fr["cols"]] = fr["ppm"]
        rgba = outputs.ramp_rgba(a, H2S_RAMP)
        src = drv.Create("", g.ncols, g.nrows, 4, gdal.GDT_Byte)
        src.SetGeoTransform(g.gt)
        src.SetProjection(g.wkt)
        for b in range(4):
            src.GetRasterBand(b + 1).WriteArray(rgba[..., b])
        dst = gdal.Warp("", src, format="MEM", dstSRS="EPSG:3857", resampleAlg="bilinear")
        p = os.path.join(folder, "%s_h2s_%03d.png" % (prefix, i))
        png.CreateCopy(p, dst)
        if os.path.exists(p + ".aux.xml"):
            os.remove(p + ".aux.xml")
        paths.append(p)
        if bounds is None:
            dgt = dst.GetGeoTransform()
            x0, y1 = dgt[0], dgt[3]
            x1 = x0 + dst.RasterXSize * dgt[1]
            y0 = y1 + dst.RasterYSize * dgt[5]
            bounds = outputs.merc_to_ll_bounds(x0, y0, x1, y1)
    log("%d H2S frames rendered." % len(paths))
    return paths, bounds


def _write_vectors(res, path, source_xy, recs, zone_m, label):
    g = res["grid"]
    if os.path.exists(path):
        os.remove(path)
    drv = ogr.GetDriverByName("GPKG")
    ds = drv.CreateDataSource(path)
    srs = osr.SpatialReference()
    srs.ImportFromWkt(g.wkt)
    lyr = ds.CreateLayer("h2s_threshold_zones", srs, ogr.wkbMultiPolygon)
    for nm, tp in (("threshold", ogr.OFTString), ("ppm", ogr.OFTReal), ("basis", ogr.OFTString),
                   ("max_dist_m", ogr.OFTReal), ("area_ha", ogr.OFTReal), ("colour", ogr.OFTString)):
        lyr.CreateField(ogr.FieldDefn(nm, tp))
    for key, lab, thr, basis, col in THRESHOLDS:
        arr = res["max1h"] if basis == "1h" else res["peak"]
        m = arr >= thr
        if not m.any():
            continue
        mem = outputs._mem_raster(m.astype(np.float32), g.gt, g.wkt)
        tmp = ogr.GetDriverByName("Memory").CreateDataSource("t")
        tl = tmp.CreateLayer("p", srs, ogr.wkbPolygon)
        tl.CreateField(ogr.FieldDefn("v", ogr.OFTInteger))
        gdal.Polygonize(mem.GetRasterBand(1), mem.GetRasterBand(1), tl, 0)
        geom = ogr.Geometry(ogr.wkbMultiPolygon)
        for f in tl:
            if f.GetField("v") == 1:
                geom.AddGeometry(f.GetGeometryRef().Clone())
        geom = geom.SimplifyPreserveTopology(g.dx * 0.6)
        geom = geom.MakeValid() if hasattr(geom, "MakeValid") else geom
        r_, c_ = np.nonzero(m)
        d = np.hypot(res["X"][c_] - source_xy[0], res["Y"][r_] - source_xy[1])
        ft = ogr.Feature(lyr.GetLayerDefn())
        ft.SetField("threshold", lab)
        ft.SetField("ppm", thr)
        ft.SetField("basis", "highest 60-minute mean" if basis == "1h" else "highest 10-minute mean")
        ft.SetField("max_dist_m", round(float(d.max()), 1))
        ft.SetField("area_ha", round(float(m.sum()) * g.cell_area / 1e4, 2))
        ft.SetField("colour", col)
        ft.SetGeometry(ogr.ForceToMultiPolygon(geom))
        lyr.CreateFeature(ft)
    # exclusion zone for comparison
    if zone_m:
        z = ds.CreateLayer("h2s_planned_zone", srs, ogr.wkbPolygon)
        z.CreateField(ogr.FieldDefn("name", ogr.OFTString))
        z.CreateField(ogr.FieldDefn("radius_m", ogr.OFTReal))
        pt = ogr.CreateGeometryFromWkt("POINT (%f %f)" % tuple(source_xy))
        ft = ogr.Feature(z.GetLayerDefn())
        ft.SetField("name", label or "Planned H2S zone")
        ft.SetField("radius_m", zone_m)
        ft.SetGeometry(pt.Buffer(zone_m, 32))
        z.CreateFeature(ft)
    s = ds.CreateLayer("h2s_source", srs, ogr.wkbPoint)
    s.CreateField(ogr.FieldDefn("name", ogr.OFTString))
    ft = ogr.Feature(s.GetLayerDefn())
    ft.SetField("name", "H2S leak point")
    ft.SetGeometry(ogr.CreateGeometryFromWkt("POINT (%f %f)" % tuple(source_xy)))
    s.CreateFeature(ft)
    rl = ds.CreateLayer("h2s_receptor_exposure", srs, ogr.wkbPoint)
    for nm, tp in (("name", ogr.OFTString), ("type", ogr.OFTString), ("dist_km", ogr.OFTReal),
                   ("hazard", ogr.OFTString), ("peak_ppm", ogr.OFTReal), ("max_1h_ppm", ogr.OFTReal),
                   ("first_aegl2_h", ogr.OFTReal), ("first_idlh_h", ogr.OFTReal)):
        rl.CreateField(ogr.FieldDefn(nm, tp))
    for r in recs:
        ft = ogr.Feature(rl.GetLayerDefn())
        for k in ("name", "type", "hazard"):
            ft.SetField(k, str(r[k]))
        ft.SetField("dist_km", round(r["dist_km"], 3))
        ft.SetField("peak_ppm", round(r["peak_ppm"], 3))
        ft.SetField("max_1h_ppm", round(r["max_1h_ppm"], 3))
        if r["first_aegl2_h"] is not None:
            ft.SetField("first_aegl2_h", r["first_aegl2_h"])
        if r["first_idlh_h"] is not None:
            ft.SetField("first_idlh_h", r["first_idlh_h"])
        ft.SetGeometry(ogr.CreateGeometryFromWkt("POINT (%f %f)" % (r["x"], r["y"])))
        rl.CreateFeature(ft)
    ds = None
    return path


def run_h2s(cfg, log=print, feedback=None):
    """cfg keys: name, well, dem, source_xy, gas_mmscfd, h2s_ppm, exit_d_m, release_h_m,
    gas_temp_c, gas_mw, duration_h, sim_hours, frame_minutes, month, start_iso, wind_csv,
    wind_seed, half_width_m, cell_m, puff_s, receptor_specs, html_layers, zone_m, zone_label,
    out_dir, prefix, prepared_by, title, subtitle, notes, html (bool), links."""
    c = dict(gas_mmscfd=10.0, h2s_ppm=10000.0, exit_d_m=0.20, release_h_m=10.0, gas_temp_c=40.0,
             gas_mw=19.0, duration_h=24.0, sim_hours=24.0, frame_minutes=10.0, month=0,
             start_iso=None, wind_csv=None, wind_seed=7, half_width_m=10000.0, cell_m=50.0,
             puff_s=90.0, receptor_specs=[], html_layers=[], zone_m=None, zone_label=None,
             out_dir=".", prefix=None, prepared_by="", notes=[], html=True, jet=True, links=[])
    c.update({k: v for k, v in cfg.items() if v is not None})
    os.makedirs(c["out_dir"], exist_ok=True)
    prefix = c["prefix"] or c["name"].replace(" ", "_")
    clim = climate.month_climate(c["month"])
    if c.get("wind_csv") and os.path.exists(c["wind_csv"]):
        u, v, s = climate.load_wind_csv(c["wind_csv"])
        clim["u"], clim["v"], clim["speed"] = u[c["month"]], v[c["month"]], s[c["month"]]
    start_hour, doy = 0, 15 + 30 * int(c["month"])
    if c.get("start_iso"):
        d0 = _dt.datetime.fromisoformat(c["start_iso"])
        start_hour, doy = d0.hour + d0.minute / 60.0, d0.timetuple().tm_yday
    hours = int(c["sim_hours"]) + 3
    wind = fallout.wind_series(clim["u"], clim["v"], clim["speed"], hours + int(start_hour),
                               seed=c["wind_seed"])[int(start_hour):]
    src = H2SSource(c["gas_mmscfd"], c["h2s_ppm"], c["exit_d_m"], c["release_h_m"], c["gas_temp_c"],
                    c["gas_mw"], c["duration_h"], jet=c["jet"], buildup_h=c.get("buildup_h", 0.0))
    # site latitude / longitude
    ds = gdal.Open(c["dem"])
    tr = html_export._to_ll(ds.GetProjection())
    ds = None
    lon, lat, _ = tr.TransformPoint(*c["source_xy"])
    res = simulate(c["dem"], c["source_xy"], src, wind, c["month"], start_hour, doy, c["sim_hours"],
                   frame_minutes=c["frame_minutes"], half_width_m=c["half_width_m"], cell_m=c["cell_m"],
                   puff_s=c["puff_s"], lat=lat, lon=lon, cloud=clim["cloud"], t_mean=clim["temp_c"],
                   t_max=clim["tmax_c"], log=log, feedback=feedback)
    g = res["grid"]
    recs = load_receptors(c["receptor_specs"], g.wkt, g.extent(), log=log)
    rx = receptor_exposure(res, recs, c["source_xy"], c["frame_minutes"] / 60.0)
    folder = c["out_dir"]
    files = {}
    # rasters
    for nm, arr, nd in (("h2s_max_10min_ppm", res["peak"], 0), ("h2s_max_60min_ppm", res["max1h"], 0)):
        p = os.path.join(folder, "%s_%s.tif" % (prefix, nm))
        outputs.write_tif(p, arr, g.gt, g.wkt, nd)
        files[nm] = p
    ta = res["t_aegl2"].copy()
    ta[np.isnan(ta)] = -1
    p = os.path.join(folder, "%s_h2s_time_to_AEGL2_h.tif" % prefix)
    outputs.write_tif(p, ta, g.gt, g.wkt, -1)
    files["t_aegl2"] = p
    stack = np.zeros((len(res["frames"]), g.nrows, g.ncols), np.float32)
    for i, fr in enumerate(res["frames"]):
        stack[i][fr["rows"], fr["cols"]] = fr["ppm"]
    p = os.path.join(folder, "%s_h2s_ppm_timeseries.tif" % prefix)
    dsx = outputs.write_tif(p, stack, g.gt, g.wkt, 0)
    for i, fr in enumerate(res["frames"]):
        dsx.GetRasterBand(i + 1).SetDescription("T+%.2f h" % fr["t_h"])
    dsx = None
    files["timeseries"] = p
    files["gpkg"] = _write_vectors(res, os.path.join(folder, prefix + "_h2s_results.gpkg"),
                                   c["source_xy"], rx, c.get("zone_m"), c.get("zone_label"))
    # tables
    p = os.path.join(folder, prefix + "_h2s_timeseries.csv")
    keys = list(res["stats"][0].keys())
    with open(p, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(keys)
        for r in res["stats"]:
            w.writerow([round(r[k], 3) if isinstance(r[k], float) else r[k] for k in keys])
    files["stats_csv"] = p
    p = os.path.join(folder, prefix + "_h2s_receptors.csv")
    rk = ["name", "type", "dist_km", "hazard", "peak_ppm", "max_1h_ppm", "first_aegl1_h",
          "first_aegl2_h", "first_aegl3_h", "first_idlh_h", "x", "y"]
    rx_sorted = sorted(rx, key=lambda r: -r["peak_ppm"])
    with open(p, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(rk)
        for r in rx_sorted:
            w.writerow(["" if r[k] is None else (round(r[k], 3) if isinstance(r[k], float) else r[k])
                        for k in rk])
    files["receptors_csv"] = p
    summary = summarise(res, c, src, rx)
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Scenario"
        for k, v in summary_rows(summary):
            ws.append([k, v])
        ws.column_dimensions["A"].width = 44
        ws.column_dimensions["B"].width = 60
        ws2 = wb.create_sheet("Time series")
        ws2.append(keys)
        for r in res["stats"]:
            ws2.append([round(r[k], 3) if isinstance(r[k], float) else r[k] for k in keys])
        ws3 = wb.create_sheet("Receptors")
        ws3.append(rk)
        for r in rx_sorted:
            ws3.append([None if r[k] is None else (round(r[k], 3) if isinstance(r[k], float) else r[k])
                        for k in rk])
        for sh in (ws2, ws3):
            for cell in sh[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E79")
        wb.properties.creator = c.get("prepared_by", "")
        p = os.path.join(folder, prefix + "_h2s_results.xlsx")
        wb.save(p)
        files["xlsx"] = p
    except ImportError:
        pass
    with open(os.path.join(folder, prefix + "_h2s_summary.json"), "w") as f:
        json.dump(summary, f, indent=1, default=float)
    fdir = os.path.join(folder, prefix + "_frames")
    pngs, bounds = _ramp_png_frames(res, fdir, prefix, log)
    files.update(frames=pngs, frame_bounds=bounds)
    # hillshade for the offline basemap
    hs = _hillshade_png(g, os.path.join(fdir, prefix + "_hillshade.png"))
    files["hillshade"], files["hill_bounds"] = hs
    rec = dict(version=1, model="h2s", name=c["name"], well=c.get("well"), title=c.get("title") or c["name"],
               subtitle=c.get("subtitle") or "H2S dispersion simulation", kind="H2S",
               start_iso=c.get("start_iso"), source_xy=list(c["source_xy"]), crs_wkt=g.wkt,
               frame_minutes=c["frame_minutes"], timeseries=os.path.basename(files["timeseries"]),
               units="ppm", ramp=H2S_RAMP, frames=[dict(t_h=fr["t_h"], rain_mmh=0.0, wind=fr["wind"],
                                                       stab=fr["stab"]) for fr in res["frames"]],
               stats=res["stats"], summary=summary, prepared_by=c.get("prepared_by", ""),
               duration_h=c["duration_h"], sim_hours=c["sim_hours"])
    p = os.path.join(folder, prefix + "_run.json")
    with open(p, "w") as f:
        json.dump(rec, f, indent=1, default=float)
    files["run_json"] = p
    files["summary"] = summary
    try:
        from .runner import add_captions
        files["captions"] = add_captions(p, files["timeseries"], rx, log=log)
    except Exception as ex:
        log("Captions not written: %s" % ex)
        files["captions"] = []
    if c.get("html", True):
        files["html"] = write_html(res, c, files, summary, rx, src, prefix, log)
    res["files"] = files
    res["receptors"] = rx
    res["summary"] = summary
    log("H2S run complete: AEGL-2 zone reaches %.0f m, IDLH %.0f m from the release." % (
        summary["dist_AEGL2_m"], summary["dist_IDLH_m"]))
    return res


def _hillshade_png(g, path):
    from .terrain import relief_shade

    class _T(object):
        pass
    t = _T()
    t.grid, t.z_raw, t.hillshade = g, g.z, relief_shade(g.z, g.dx)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    return path, outputs.hillshade_png(t, None, path)


def summarise(res, c, src, rx):
    g = res["grid"]
    out = dict(name=c["name"], well=c.get("well"), month=climate.MONTHS[int(c["month"])],
               gas_mmscfd=src.gas_mmscfd, h2s_ppm=src.h2s_ppm, h2s_gs=src.q_gs,
               h2s_kg_total=src.q_gs * c["duration_h"] * 3.6, exit_velocity=src.v_exit,
               release_h_m=src.release_h_m, duration_h=c["duration_h"], sim_hours=c["sim_hours"],
               buildup_h=src.buildup_h,
               dense_gas=bool(res["dense"]), max_ppm=float(res["peak"].max()),
               zone_m=c.get("zone_m"))
    for key, lab, thr, basis, col in THRESHOLDS:
        arr = res["max1h"] if basis == "1h" else res["peak"]
        m = arr >= thr
        if m.any():
            r_, c_ = np.nonzero(m)
            d = np.hypot(res["X"][c_] - c["source_xy"][0], res["Y"][r_] - c["source_xy"][1])
            out["dist_%s_m" % key] = float(d.max())
            out["area_%s_ha" % key] = float(m.sum()) * g.cell_area / 1e4
        else:
            out["dist_%s_m" % key] = 0.0
            out["area_%s_ha" % key] = 0.0
    out["stability_counts"] = {k: sum(1 for s in res["stats"] if s["stability"] == k) for k in "ABCDEF"}
    out["receptors_aegl2"] = [dict(name=r["name"], type=r["type"], dist_km=r["dist_km"],
                                   first_h=r["first_aegl2_h"], peak_ppm=r["peak_ppm"])
                              for r in sorted(rx, key=lambda r: r["dist_km"]) if r["max_1h_ppm"] >= 27.0]
    out["receptors_aegl1"] = [dict(name=r["name"], type=r["type"], dist_km=r["dist_km"],
                                   max_1h_ppm=r["max_1h_ppm"])
                              for r in sorted(rx, key=lambda r: r["dist_km"])
                              if 0.51 <= r["max_1h_ppm"] < 27.0]
    return out


def summary_rows(s):
    rows = [("Scenario", s["name"]), ("Well", s["well"]), ("Month", s["month"]),
            ("Gas rate (MMscf/d)", s["gas_mmscfd"]), ("H2S in gas (ppm mol)", s["h2s_ppm"]),
            ("H2S leak rate (g/s)", round(s["h2s_gs"], 3)),
            ("H2S leaked over the leak period (kg)", round(s["h2s_kg_total"], 1)),
            ("Release height (m) / exit velocity (m/s)", "%.0f / %.0f" % (s["release_h_m"], s["exit_velocity"])),
            ("Release duration (h) / simulated (h)", "%.0f / %.0f" % (s["duration_h"], s["sim_hours"])),
            ("Leak build-up", ("from 10 %% to the full rate over %.1f h" % s["buildup_h"]) if s.get("buildup_h")
             else "full rate from the start (worst case)"),
            ("Highest ground-level 10-min concentration (ppm)", round(s["max_ppm"], 2))]
    for key, lab, thr, basis, col in THRESHOLDS:
        rows.append(("Furthest extent, %s (m)" % lab, round(s["dist_%s_m" % key], 0)))
    if s.get("zone_m"):
        rows.append(("Planned H2S zone radius for comparison (m)", s["zone_m"]))
    rows.append(("Dense-gas warning", "yes" if s["dense_gas"] else "no (buoyant)"))
    return rows


def write_html(res, c, files, summary, rx, src, prefix, log=print):
    g = res["grid"]
    tr = html_export._to_ll(g.wkt)
    lon, lat, _ = tr.TransformPoint(*c["source_xy"])
    caps = files.get("captions") or []
    frames = [dict(t=round(fr["t_h"], 3), img=html_export._uri(files["frames"][i]),
                   cap=caps[i] if i < len(caps) else None,
                   wind=[round(fr["wind"][0], 1), round(fr["wind"][1], 2)], stab=fr["stab"])
              for i, fr in enumerate(res["frames"])]
    hb = files["hill_bounds"]
    clip = [hb[0] - 0.01, hb[1] - 0.01, hb[2] + 0.01, hb[3] + 0.01]
    layers = []
    zl = [dict(name="H2S threshold zones (worst over the run)", path=files["gpkg"], layer="h2s_threshold_zones",
               kind="zones", label="threshold", visible=True),
          dict(name=c.get("zone_label") or "Planned H2S zone", path=files["gpkg"], layer="h2s_planned_zone",
               style=dict(color="#000000", weight=2, dash="6 4", fillOpacity=0), visible=True, label="name")]
    for L in zl + list(c["html_layers"]):
        gj = html_export.layer_geojson(L["path"], L.get("layer"), clip_ll=clip, simplify_m=L.get("simplify", 0.0),
                                       where=L.get("where"))
        if not gj or not gj["features"]:
            continue
        st = dict(L.get("style", {}))
        st["label"] = L.get("label", "name")
        if L.get("permanent"):
            st["permanent"] = True
        layers.append(dict(name=L["name"], kind=L.get("kind", "vector"), geojson=gj, style=st,
                           visible=L.get("visible", True)))
    recs = []
    for r in rx:
        rl, rb, _ = tr.TransformPoint(r["x"], r["y"])
        recs.append(dict(name=r["name"], type=r["type"], lat=round(rb, 6), lon=round(rl, 6),
                         dist=round(r["dist_km"], 3), peak=round(r["peak_ppm"], 3),
                         m1=round(r["max_1h_ppm"], 3), hazard=r["hazard"],
                         a2=r["first_aegl2_h"], a1=r["first_aegl1_h"], idlh=r["first_idlh_h"],
                         series=[round(v, 2) for v in r["series"]] if r["peak_ppm"] >= 0.05 else None))
    facts = [[k, str(v)] for k, v in summary_rows(summary)] + [["Prepared by", c.get("prepared_by") or "-"]]
    notes = list(c.get("notes") or []) + [
        "Gaussian puff model: puffs released every %.0f s, carried by the hourly wind at plume height; "
        "Briggs plume rise; Pasquill-Gifford stability by the Turner scheme; Briggs open-country spreads; "
        "reflection at the ground and at the mixing height; half-height terrain correction." % c["puff_s"],
        "Concentrations are ground-level 10-minute means; AEGL zones use the highest 60-minute running mean, "
        "IDLH and 500 ppm zones the highest 10-minute mean.",
        "Not modelled: ignition of the release (an ignited leak converts H2S to SO2), dense-gas slumping "
        "of H2S-rich mixtures, building or rig-structure wake effects, calm-wind stagnation below 1 m/s."]
    stats = [{k: (round(v, 2) if isinstance(v, float) else v) for k, v in s.items()} for s in res["stats"]]
    payload = dict(
        meta=dict(title=c.get("title") or c["name"], subtitle=c.get("subtitle") or "H2S dispersion simulation",
                  source=[lat, lon], source_label=c.get("source_label") or "Leak point",
                  start_iso=c.get("start_iso"), facts=facts, notes=notes, frame_h=c["frame_minutes"] / 60.0,
                  q_gs=round(src.q_gs, 3), duration_h=c["duration_h"], zone_m=c.get("zone_m"),
                  links=c.get("links") or []),
        frames=frames, bounds=files["frame_bounds"], hill=dict(uri=html_export._uri(files["hillshade"]), bounds=hb),
        stats=stats, receptors=recs, layers=layers, ramp=H2S_RAMP,
        thresholds=[dict(key=k, label=lab, ppm=t, basis=b, color=col) for k, lab, t, b, col in THRESHOLDS])
    p = os.path.join(c["out_dir"], prefix + "_h2s_animation.html")
    html_export.build_h2s_html(p, payload)
    log("H2S HTML written: %s (%.1f MB)" % (p, os.path.getsize(p) / 1e6))
    return p
