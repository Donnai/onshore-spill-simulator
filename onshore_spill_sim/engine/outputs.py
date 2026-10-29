# -*- coding: utf-8 -*-
"""
Writers for the simulation results: GeoTIFFs, GeoPackage layers, tables,
colour-mapped PNG frames (EPSG:3857, for web maps) and the summary record.
"""
import csv
import json
import math
import os
import numpy as np
from osgeo import gdal, ogr, osr

gdal.UseExceptions()

# "rain radar" ramp for surface oil thickness (mm), log steps
OIL_RAMP = [  # (mm, r, g, b, a) - yellow to purple, clear of water blues and savannah greens
    (0.01, 255, 247, 150, 120),
    (0.1, 255, 225, 60, 185),
    (0.5, 255, 170, 20, 215),
    (1.0, 250, 120, 10, 230),
    (2.5, 235, 60, 20, 240),
    (5.0, 200, 20, 60, 245),
    (10.0, 160, 0, 110, 250),
    (25.0, 105, 0, 135, 252),
    (50.0, 60, 10, 90, 255),
    (100.0, 30, 10, 40, 255),
]
# oil temperature (C), linear
TEMP_RAMP = [(15.0, 49, 54, 149, 200), (22.0, 69, 117, 180, 210), (28.0, 171, 217, 233, 220),
             (34.0, 254, 224, 144, 230), (42.0, 244, 109, 67, 240), (55.0, 165, 0, 38, 250)]
# yield stress (Pa), log steps
YIELD_RAMP = [(0.1, 230, 245, 255, 90), (1.0, 158, 202, 225, 160), (10.0, 66, 146, 198, 200),
              (100.0, 8, 81, 156, 230), (1000.0, 8, 29, 88, 250)]
# well spray fallout ("oil rain") deposit, mm
FALL_RAMP = [
    (0.005, 125, 105, 80, 25),
    (0.02, 110, 85, 60, 60),
    (0.1, 90, 62, 38, 110),
    (1.0, 62, 40, 22, 170),
    (5.0, 30, 20, 10, 215),
]


def ramp_rgba(values_mm, ramp):
    v = np.asarray(values_mm, dtype=np.float64)
    out = np.zeros(v.shape + (4,), dtype=np.uint8)
    lv = np.log10(np.maximum(v, 1e-9))
    stops = np.log10([s[0] for s in ramp])
    cols = np.array([s[1:] for s in ramp], dtype=np.float64)
    on = v >= ramp[0][0]
    for ch in range(4):
        out[..., ch] = np.where(on, np.interp(lv, stops, cols[:, ch]), 0).astype(np.uint8)
    return out


def ramp_rgba_lin(values, ramp, valid=None):
    """Linear colour ramp (for temperature); valid = boolean mask of cells to colour."""
    v = np.asarray(values, dtype=np.float64)
    out = np.zeros(v.shape + (4,), dtype=np.uint8)
    stops = np.array([s[0] for s in ramp], dtype=np.float64)
    cols = np.array([s[1:] for s in ramp], dtype=np.float64)
    on = np.ones(v.shape, bool) if valid is None else valid
    for ch in range(4):
        out[..., ch] = np.where(on, np.interp(v, stops, cols[:, ch]), 0).astype(np.uint8)
    return out


def field_pngs_3857(result, window, folder, prefix, key, ramp, linear=False, scale=1.0, min_h=1e-5):
    """PNG frames of a per-cell oil property (temp, visc, tau) in EPSG:3857."""
    g = result.grid
    r0, r1, c0, c1 = window
    gt = sub_gt(g.gt, r0, c0)
    os.makedirs(folder, exist_ok=True)
    drv = gdal.GetDriverByName("MEM")
    png = gdal.GetDriverByName("PNG")
    paths = []
    for i, fr in enumerate(result.frames):
        a = np.zeros((g.nrows, g.ncols), np.float32)
        m = np.zeros((g.nrows, g.ncols), bool)
        ok = fr["h"] >= min_h
        a[fr["rows"][ok], fr["cols"][ok]] = fr[key][ok] * scale
        m[fr["rows"][ok], fr["cols"][ok]] = True
        a, m = a[r0:r1, c0:c1], m[r0:r1, c0:c1]
        rgba = ramp_rgba_lin(a, ramp, m) if linear else ramp_rgba(np.where(m, a, 0), ramp)
        src = drv.Create("", c1 - c0, r1 - r0, 4, gdal.GDT_Byte)
        src.SetGeoTransform(gt)
        src.SetProjection(g.wkt)
        for b in range(4):
            src.GetRasterBand(b + 1).WriteArray(rgba[..., b])
        dst = gdal.Warp("", src, format="MEM", dstSRS="EPSG:3857", resampleAlg="near")
        p = os.path.join(folder, "%s_%s_%03d.png" % (prefix, key, i))
        png.CreateCopy(p, dst)
        if os.path.exists(p + ".aux.xml"):
            os.remove(p + ".aux.xml")
        paths.append(p)
    return paths


def _srs(wkt):
    s = osr.SpatialReference()
    s.ImportFromWkt(wkt)
    s.SetAxisMappingStrategy(osr.OAMS_TRADITIONAL_GIS_ORDER)
    return s


def write_tif(path, arr, gt, wkt, nodata=None, dtype=gdal.GDT_Float32):
    drv = gdal.GetDriverByName("GTiff")
    a = np.atleast_3d(arr)
    if a.ndim == 3 and arr.ndim == 2:
        a = arr[np.newaxis, ...]
    elif arr.ndim == 3:
        a = arr
    nb, nr, nc = a.shape
    ds = drv.Create(path, nc, nr, nb, dtype, options=["COMPRESS=DEFLATE", "TILED=YES",
                                                      "PREDICTOR=1", "BIGTIFF=IF_SAFER"])
    ds.SetGeoTransform(gt)
    ds.SetProjection(wkt)
    for b in range(nb):
        band = ds.GetRasterBand(b + 1)
        if nodata is not None:
            band.SetNoDataValue(nodata)
        band.WriteArray(a[b])
    ds.FlushCache()
    return ds


def footprint_window(result, margin=8):
    """Row/col window that contains every oiled / fallout cell of the run."""
    g = result.grid
    m = (result.hmax >= 1e-6) | (result.fallout >= 1e-6) | (result.infiltrated > 1e-6)
    r, c = result.grid.xy_to_rc(*result.scenario.source_xy)
    if not m.any():
        rr, cc = np.array([r]), np.array([c])
    else:
        rr, cc = np.nonzero(m)
    r0 = max(int(rr.min()) - margin, 0)
    r1 = min(int(rr.max()) + margin + 1, g.nrows)
    c0 = max(int(cc.min()) - margin, 0)
    c1 = min(int(cc.max()) + margin + 1, g.ncols)
    # keep it at least ~2 km square around the source for context
    half = int(1000 / g.dx)
    r0, r1 = max(min(r0, r - half), 0), min(max(r1, r + half), g.nrows)
    c0, c1 = max(min(c0, c - half), 0), min(max(c1, c + half), g.ncols)
    return r0, r1, c0, c1


def frame_array(frame, shape):
    h = np.zeros(shape, dtype=np.float32)
    h[frame["rows"], frame["cols"]] = frame["h"]
    f = np.zeros(shape, dtype=np.float32)
    f[frame["frows"], frame["fcols"]] = frame["fall"]
    return h, f


def sub_gt(gt, r0, c0):
    return (gt[0] + c0 * gt[1], gt[1], 0.0, gt[3] + r0 * gt[5], 0.0, gt[5])


def write_rasters(result, folder, prefix, log=print):
    g = result.grid
    r0, r1, c0, c1 = footprint_window(result)
    gt = sub_gt(g.gt, r0, c0)
    sl = (slice(r0, r1), slice(c0, c1))
    files = {}

    def w(name, arr, nodata=None):
        p = os.path.join(folder, "%s_%s.tif" % (prefix, name))
        write_tif(p, arr.astype(np.float32), gt, g.wkt, nodata)
        files[name] = p

    w("max_thickness_mm", result.hmax[sl] * 1000.0, 0)
    ta = result.t_arrival[sl].copy()
    ta[np.isnan(ta)] = -1
    w("arrival_time_h", ta, -1)
    w("infiltrated_L_per_m2", result.infiltrated[sl] * 1000.0, 0)
    pen = result.penetration[sl].copy()
    pen[~np.isfinite(pen)] = 0
    w("penetration_depth_m", pen, 0)
    if result.fallout.any():
        w("fallout_mm", result.fallout[sl] * 1000.0, 0)
    w("final_thickness_mm", result.final_h[sl] * 1000.0, 0)
    heat = getattr(result.scenario, "heat_mode", False)
    if heat:
        for nm, arr in (("oil_temp_max_C", result.temp_max), ("oil_temp_min_C", result.temp_min)):
            a = arr[sl].copy()
            a[~np.isfinite(a)] = -999
            w(nm, a, -999)
        w("viscosity_max_mPas", result.visc_max[sl] * 1000.0, 0)
        if result.yield_max[sl].max() > 0:
            w("yield_stress_max_Pa", result.yield_max[sl], 0)
    # time series: one band per frame (surface oil, cumulative spray fallout)
    shape = (g.nrows, g.ncols)
    stack = np.zeros((len(result.frames), r1 - r0, c1 - c0), dtype=np.float32)
    fstack = np.zeros_like(stack)
    for i, fr in enumerate(result.frames):
        h, f = frame_array(fr, shape)
        stack[i] = h[sl] * 1000.0
        fstack[i] = f[sl] * 1000.0
    series = [("timeseries", stack, "thickness_timeseries_mm", 0),
              ("fallout_timeseries", fstack, "fallout_timeseries_mm", 0)]
    if heat:
        for key, fk, scale, nd in (("temp_timeseries", "temp", 1.0, -999),
                                   ("visc_timeseries", "visc", 1000.0, 0),
                                   ("yield_timeseries", "tau", 1.0, 0)):
            st = np.full_like(stack, nd)
            for i, fr in enumerate(result.frames):
                a = np.full(shape, nd, dtype=np.float32)
                a[fr["rows"], fr["cols"]] = fr[fk] * scale
                st[i] = a[sl]
            nm = {"temp": "oil_temperature_timeseries_C", "visc": "viscosity_timeseries_mPas",
                  "tau": "yield_stress_timeseries_Pa"}[fk]
            if fk == "tau" and not (st > 0).any():
                continue          # oil never gels - no yield-stress series
            series.append((key, st, nm, nd))
    for key, arr, nm, nd in series:
        if key == "fallout_timeseries" and not fstack.any():
            continue
        p = os.path.join(folder, "%s_%s.tif" % (prefix, nm))
        ds = write_tif(p, arr, gt, g.wkt, nd)
        for i, fr in enumerate(result.frames):
            ds.GetRasterBand(i + 1).SetDescription("T+%.2f h" % fr["t_h"])
        ds = None
        files[key] = p
    log("Rasters written (%d x %d cells, %d time steps)." % (c1 - c0, r1 - r0, len(result.frames)))
    return files, (r0, r1, c0, c1)


def _mem_raster(arr, gt, wkt):
    drv = gdal.GetDriverByName("MEM")
    ds = drv.Create("", arr.shape[1], arr.shape[0], 1, gdal.GDT_Float32)
    ds.SetGeoTransform(gt)
    ds.SetProjection(wkt)
    ds.GetRasterBand(1).WriteArray(arr.astype(np.float32))
    return ds


def polygonize(arr_bool, gt, wkt, layer, attrs):
    """Add polygons of True cells to layer with constant attributes."""
    ds = _mem_raster(arr_bool.astype(np.float32), gt, wkt)
    band = ds.GetRasterBand(1)
    mem = ogr.GetDriverByName("Memory").CreateDataSource("p")
    ml = mem.CreateLayer("p", _srs(wkt), ogr.wkbPolygon)
    ml.CreateField(ogr.FieldDefn("v", ogr.OFTInteger))
    gdal.Polygonize(band, band, ml, 0)
    geoms = ogr.Geometry(ogr.wkbMultiPolygon)
    for f in ml:
        if f.GetField("v") == 1:
            geoms.AddGeometry(f.GetGeometryRef().Clone())
    if geoms.GetGeometryCount() == 0:
        return 0.0
    u = geoms.UnionCascaded()
    u = ogr.ForceToMultiPolygon(u)
    feat = ogr.Feature(layer.GetLayerDefn())
    for k, v in attrs.items():
        feat.SetField(k, v)
    feat.SetField("area_ha", u.GetArea() / 1e4)
    feat.SetGeometry(u)
    layer.CreateFeature(feat)
    return u.GetArea() / 1e4


def write_vectors(result, path, prefix, receptors, streams=None, log=print):
    """GeoPackage with footprints, isochrones, receptor results and streams."""
    g = result.grid
    if os.path.exists(path):
        gdal.GetDriverByName("GPKG").Delete(path)
    ds = ogr.GetDriverByName("GPKG").CreateDataSource(path)
    srs = _srs(g.wkt)
    sc = result.scenario
    # footprints by thickness class
    l = ds.CreateLayer("oil_footprint", srs, ogr.wkbMultiPolygon)
    for f, t in (("class", ogr.OFTString), ("min_mm", ogr.OFTReal), ("area_ha", ogr.OFTReal)):
        l.CreateField(ogr.FieldDefn(f, t))
    for mm, lab in ((0.1, "Oiled ground >= 0.1 mm"), (1.0, ">= 1 mm"), (10.0, ">= 10 mm (pooled oil)")):
        polygonize(result.hmax * 1000 >= mm, g.gt, g.wkt, l, {"class": lab, "min_mm": mm})
    # isochrones
    l = ds.CreateLayer("arrival_isochrones", srs, ogr.wkbMultiPolygon)
    for f, t in (("label", ogr.OFTString), ("hours", ogr.OFTReal), ("area_ha", ogr.OFTReal)):
        l.CreateField(ogr.FieldDefn(f, t))
    ta = np.nan_to_num(result.t_arrival, nan=1e9)
    for hh in (1, 3, 6, 12, 24, 48, 72, 96, 120, 168):
        if hh > sc.sim_hours + 1e-6:
            break
        polygonize(ta <= hh, g.gt, g.wkt, l, {"label": "Oil within %d h" % hh, "hours": hh})
    # well spray fallout
    l = ds.CreateLayer("spray_fallout", srs, ogr.wkbMultiPolygon)
    for f, t in (("class", ogr.OFTString), ("min_mm", ogr.OFTReal), ("area_ha", ogr.OFTReal)):
        l.CreateField(ogr.FieldDefn(f, t))
    for mm, lab in ((0.01, "Light oil mist >= 0.01 mm"), (0.1, "Oil spotting >= 0.1 mm"),
                    (1.0, "Heavy deposit >= 1 mm")):
        polygonize(result.fallout * 1000 >= mm, g.gt, g.wkt, l, {"class": lab, "min_mm": mm})
    # groundwater concern
    l = ds.CreateLayer("groundwater_concern", srs, ogr.wkbMultiPolygon)
    for f, t in (("class", ogr.OFTString), ("min_mm", ogr.OFTReal), ("area_ha", ogr.OFTReal)):
        l.CreateField(ogr.FieldDefn(f, t))
    pen = np.nan_to_num(result.penetration, nan=0, posinf=0)
    polygonize(pen >= sc.dtw_m, g.gt, g.wkt, l, {"class": "Oil front reaches the water table (%.1f m)" % sc.dtw_m,
                                                  "min_mm": 0})
    polygonize((pen >= 0.5 * sc.dtw_m) & (pen < sc.dtw_m), g.gt, g.wkt, l,
               {"class": "Oil front within half the depth to water", "min_mm": 0})
    # receptors
    l = ds.CreateLayer("receptor_impacts", srs, ogr.wkbPoint)
    fields = [("name", ogr.OFTString), ("type", ogr.OFTString), ("dist_km", ogr.OFTReal),
              ("status", ogr.OFTString), ("arrival_h", ogr.OFTReal), ("max_mm", ogr.OFTReal),
              ("fallout_mm", ogr.OFTReal), ("infil_mm", ogr.OFTReal), ("pen_m", ogr.OFTReal),
              ("gw_reached", ogr.OFTString)]
    for f, t in fields:
        l.CreateField(ogr.FieldDefn(f, t))
    for r in receptors:
        f = ogr.Feature(l.GetLayerDefn())
        f.SetField("name", r["name"])
        f.SetField("type", r["type"])
        f.SetField("dist_km", round(r["dist_km"], 3))
        f.SetField("status", receptor_status(r))
        if r["arrival_h"] is not None:
            f.SetField("arrival_h", round(r["arrival_h"], 2))
        f.SetField("max_mm", round(r["max_mm"], 3))
        f.SetField("fallout_mm", round(r["fallout_mm"], 4))
        f.SetField("infil_mm", round(r["infil_mm"], 3))
        f.SetField("pen_m", round(r["pen_m"], 3))
        f.SetField("gw_reached", "Yes" if r["gw"] else "No")
        f.SetGeometry(ogr.CreateGeometryFromWkt("POINT (%f %f)" % (r["x"], r["y"])))
        l.CreateFeature(f)
    # source
    l = ds.CreateLayer("spill_source", srs, ogr.wkbPoint)
    for f, t in (("scenario", ogr.OFTString), ("kind", ogr.OFTString), ("product", ogr.OFTString)):
        l.CreateField(ogr.FieldDefn(f, t))
    f = ogr.Feature(l.GetLayerDefn())
    f.SetField("scenario", sc.name)
    f.SetField("kind", sc.kind)
    f.SetField("product", sc.oil["label"])
    f.SetGeometry(ogr.CreateGeometryFromWkt("POINT (%f %f)" % sc.source_xy))
    l.CreateFeature(f)
    if streams:
        l = ds.CreateLayer("drainage_network_dem", srs, ogr.wkbLineString)
        l.CreateField(ogr.FieldDefn("area_km2", ogr.OFTReal))
        l.CreateField(ogr.FieldDefn("order", ogr.OFTString))
        for pts, a in streams:
            ln = ogr.Geometry(ogr.wkbLineString)
            for x, y in pts:
                ln.AddPoint_2D(x, y)
            f = ogr.Feature(l.GetLayerDefn())
            f.SetField("area_km2", round(a, 3))
            f.SetField("order", "major" if a >= 50 else ("stream" if a >= 5 else "minor"))
            f.SetGeometry(ln)
            l.CreateFeature(f)
    ds = None
    log("Vector results written: %s" % path)
    return path


def receptor_status(r):
    if r["gw"]:
        return "Oil reaches groundwater"
    if r["arrival_h"] is not None:
        return "Oiled"
    if r["fallout_mm"] >= 0.001:
        return "Spray fallout only"
    if not r.get("inside", True):
        return "Outside model domain"
    return "Not reached"


def write_tables(result, folder, prefix, receptors, log=print):
    files = {}
    p = os.path.join(folder, "%s_mass_balance.csv" % prefix)
    keys = ["t_h", "released", "contained", "surface", "on_water", "stranded", "infiltrated",
            "evaporated", "spray_out", "exported", "recovered", "area_ha", "footprint_ha",
            "rain_mmh", "rain_cum_mm", "t_air", "solar_wm2", "oil_t_mean", "gelled_m3",
            "balance_err"]

    def _v(x):
        return None if (x is None or (isinstance(x, float) and not math.isfinite(x))) else round(x, 4)
    with open(p, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(keys)
        for row in result.mass_balance:
            w.writerow(["" if _v(row.get(k)) is None else _v(row.get(k)) for k in keys])
    files["mass_balance_csv"] = p
    p = os.path.join(folder, "%s_receptors.csv" % prefix)
    with open(p, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["name", "type", "dist_km", "status", "arrival_h", "max_mm", "fallout_mm",
                    "infiltrated_mm", "penetration_m", "groundwater_reached", "x", "y"])
        for r in sorted(receptors, key=lambda r: (r["arrival_h"] is None, r["arrival_h"] or 0,
                                                  r["dist_km"])):
            w.writerow([r["name"], r["type"], round(r["dist_km"], 3), receptor_status(r),
                        "" if r["arrival_h"] is None else round(r["arrival_h"], 2),
                        round(r["max_mm"], 3), round(r["fallout_mm"], 4), round(r["infil_mm"], 3),
                        round(r["pen_m"], 3), "Yes" if r["gw"] else "No",
                        round(r["x"], 2), round(r["y"], 2)])
    files["receptors_csv"] = p
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Scenario"
        hdr = Font(bold=True, color="FFFFFF")
        fill = PatternFill("solid", fgColor="1F4E79")
        for k, v in scenario_rows(result):
            ws.append([k, v])
        ws.column_dimensions["A"].width = 38
        ws.column_dimensions["B"].width = 70
        ws2 = wb.create_sheet("Mass balance (m3)")
        ws2.append(keys)
        for row in result.mass_balance:
            ws2.append([_v(row.get(k)) for k in keys])
        ws3 = wb.create_sheet("Receptors")
        ws3.append(["Name", "Type", "Distance (km)", "Status", "Arrival (h)", "Max thickness (mm)",
                    "Spray fallout (mm)", "Infiltrated (L/m2)", "Penetration (m)", "Groundwater"])
        for r in sorted(receptors, key=lambda r: (r["arrival_h"] is None, r["arrival_h"] or 0,
                                                  r["dist_km"])):
            ws3.append([r["name"], r["type"], round(r["dist_km"], 3), receptor_status(r),
                        None if r["arrival_h"] is None else round(r["arrival_h"], 2),
                        round(r["max_mm"], 3), round(r["fallout_mm"], 4), round(r["infil_mm"], 3),
                        round(r["pen_m"], 3), "Yes" if r["gw"] else "No"])
        for sh in (ws2, ws3):
            for c in sh[1]:
                c.font = hdr
                c.fill = fill
                c.alignment = Alignment(wrap_text=True)
            for col in sh.columns:
                sh.column_dimensions[col[0].column_letter].width = 16
        p = os.path.join(folder, "%s_results.xlsx" % prefix)
        wb.properties.creator = result.scenario.prepared_by if hasattr(result.scenario, "prepared_by") else ""
        wb.save(p)
        files["xlsx"] = p
    except ImportError:
        log("openpyxl not available - Excel workbook skipped (CSV tables written).")
    return files


def scenario_rows(result):
    sc = result.scenario
    mb = result.mass_balance[-1]
    rows = [("Scenario", sc.name), ("Release type", sc.kind), ("Product", sc.oil["label"]),
            ("Source E / N (EPSG:32630)", "%.2f / %.2f" % sc.source_xy)]
    if sc.kind == "TANK":
        rows += [("Tank inventory (m3)", sc.volume_m3), ("Release time (min)", sc.release_minutes)]
    else:
        rows += [("Release rate (bbl/d)", sc.rate_bbl_d), ("Release duration (h)", sc.duration_h),
                 ("Spray fraction", sc.spray_fraction), ("Effective release height (m)", sc.release_height_m),
                 ("Droplet d50 (um)", sc.d50_um)]
    rows += [("Held in cellar / bund (m3)", sc.contained_m3), ("Simulated period (h)", sc.sim_hours),
             ("Month", getattr(sc, "month_name", "")), ("Mean air temperature (C)", sc.temp_c),
             ("Heat-transfer mode", "on" if sc.heat_mode else "off (constant properties)"),
             ("Storms (start h, duration h, depth mm)", "; ".join("%.0f, %.0f, %.0f" % r for r in sc.rain) or "none"),
             ("Stream specific discharge (L/s/km2)", sc.q_spec),
             ("Soil water saturation", sc.soil_sat), ("Depth to groundwater (m)", sc.dtw_m),
             ("Released (m3)", round(mb["released"], 2)),
             ("Evaporated (m3)", round(mb["evaporated"], 2)),
             ("Infiltrated (m3)", round(mb["infiltrated"], 2)),
             ("On the ground surface (m3)", round(mb["surface"], 2)),
             ("Stranded on stream banks (m3)", round(mb["stranded"], 2)),
             ("On open water (m3)", round(mb["on_water"], 2)),
             ("Left the model domain (m3)", round(mb["exported"] + mb["spray_out"], 2)),
             ("Recovered at sumps / recovery points (m3)", round(mb.get("recovered", 0.0), 2)),
             ("Oiled footprint (ha)", round(mb["footprint_ha"], 2))]
    if sc.heat_mode:
        rows += [("Source oil temperature (C)", sc.discharge_temp_c if sc.discharge_temp_c is not None else "air"),
                 ("Daily maximum air temperature (C)", sc.t_max_c), ("Relative humidity (%)", sc.rh_pct),
                 ("Cloud / haze fraction", sc.cloud), ("Pour point (C)", sc.oil["pour_point_c"]),
                 ("Lowest / highest oil temperature (C)",
                  "%.1f / %.1f" % (np.nanmin(result.temp_min), np.nanmax(result.temp_max))
                  if np.isfinite(result.temp_min).any() else "-"),
                 ("Highest yield stress reached (Pa)", round(float(result.yield_max.max()), 2))]
    fc = result.first_channel
    rows.append(("First entry into the drainage network",
                 "none in the simulated period" if not fc else
                 "T+%.1f h at E %.0f N %.0f (catchment %.2f km2)" % (fc["t_h"], fc["x"], fc["y"], fc["area_km2"])))
    return rows


def frame_pngs_3857(result, window, folder, prefix, log=print):
    """Colour-mapped RGBA PNGs of every frame, warped to Web Mercator.
    Returns (list of png paths, [south, west, north, east] in WGS84, fallout png list)."""
    g = result.grid
    r0, r1, c0, c1 = window
    gt = sub_gt(g.gt, r0, c0)
    shape = (g.nrows, g.ncols)
    os.makedirs(folder, exist_ok=True)
    drv = gdal.GetDriverByName("MEM")
    png = gdal.GetDriverByName("PNG")
    paths, fpaths = [], []
    bounds = None
    for i, fr in enumerate(result.frames):
        h, f = frame_array(fr, shape)
        for kind, arr, ramp, lst in (("oil", h[r0:r1, c0:c1] * 1000.0, OIL_RAMP, paths),
                                     ("fall", f[r0:r1, c0:c1] * 1000.0, FALL_RAMP, fpaths)):
            rgba = ramp_rgba(arr, ramp)
            src = drv.Create("", c1 - c0, r1 - r0, 4, gdal.GDT_Byte)
            src.SetGeoTransform(gt)
            src.SetProjection(g.wkt)
            for b in range(4):
                src.GetRasterBand(b + 1).WriteArray(rgba[..., b])
            src.GetRasterBand(4).SetColorInterpretation(gdal.GCI_AlphaBand)
            dst = gdal.Warp("", src, format="MEM", dstSRS="EPSG:3857", resampleAlg="near",
                            dstAlpha=False)
            p = os.path.join(folder, "%s_%s_%03d.png" % (prefix, kind, i))
            png.CreateCopy(p, dst)
            if os.path.exists(p + ".aux.xml"):
                os.remove(p + ".aux.xml")
            lst.append(p)
            if bounds is None:
                dgt = dst.GetGeoTransform()
                x0, y1 = dgt[0], dgt[3]
                x1 = x0 + dst.RasterXSize * dgt[1]
                y0 = y1 + dst.RasterYSize * dgt[5]
                bounds = merc_to_ll_bounds(x0, y0, x1, y1)
    log("%d animation frames rendered." % len(paths))
    return paths, bounds, fpaths


def merc_to_ll_bounds(x0, y0, x1, y1):
    R = 6378137.0
    lon = lambda x: math.degrees(x / R)
    lat = lambda y: math.degrees(2 * math.atan(math.exp(y / R)) - math.pi / 2)
    return [lat(y0), lon(x0), lat(y1), lon(x1)]


def hillshade_png(terrain, window, path):
    """Hypsometric hillshade of the whole model domain (EPSG:3857, RGBA)
    for offline display. window is ignored (kept for compatibility)."""
    g = terrain.grid
    hs = terrain.hillshade
    z = terrain.z_raw
    lo_, hi_ = np.percentile(z, 1), np.percentile(z, 99)
    zn = np.clip((z - lo_) / max(hi_ - lo_, 1e-6), 0, 1)
    # hypsometric tints (low green valleys to buff interfluves) under the relief shading
    stops = np.array([0.0, 0.35, 0.7, 1.0])
    cols = np.array([[150, 190, 120], [196, 210, 150], [226, 214, 166], [214, 186, 140]], float)
    rgb = np.stack([np.interp(zn, stops, cols[:, k]) for k in range(3)], -1)
    rgb = rgb * (0.38 + 0.72 * hs[..., None])
    rgb = np.clip(rgb, 0, 255).astype(np.uint8)
    drv = gdal.GetDriverByName("MEM")
    src = drv.Create("", g.ncols, g.nrows, 4, gdal.GDT_Byte)
    src.SetGeoTransform(g.gt)
    src.SetProjection(g.wkt)
    for b in range(3):
        src.GetRasterBand(b + 1).WriteArray(rgb[..., b])
    src.GetRasterBand(4).WriteArray(np.full((g.nrows, g.ncols), 255, np.uint8))
    src.GetRasterBand(4).SetColorInterpretation(gdal.GCI_AlphaBand)
    dst = gdal.Warp("", src, format="MEM", dstSRS="EPSG:3857", resampleAlg="bilinear")
    gdal.GetDriverByName("PNG").CreateCopy(path, dst)
    if os.path.exists(path + ".aux.xml"):
        os.remove(path + ".aux.xml")
    dgt = dst.GetGeoTransform()
    x0, y1 = dgt[0], dgt[3]
    x1 = x0 + dst.RasterXSize * dgt[1]
    y0 = y1 + dst.RasterYSize * dgt[5]
    return merc_to_ll_bounds(x0, y0, x1, y1)
